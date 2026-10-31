"""
CampusFind - Professional Simple 6-Slide Presentation Generator
==============================================================
Generates a crisp, executive, highly professional 6-slide deck:
- Clean 16:9 widescreen layout
- Professional academic palette: Deep Navy (RGB 20, 35, 60), Symbiosis Crimson (RGB 163, 29, 35), Slate Dark
- High readability, balanced whitespace, clear cards, zero clutter
"""

import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_simple_deck(output_filename="CampusFind_6Slides.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Professional Theme Colors
    C_NAVY = RGBColor(20, 35, 60)         # Primary Dark Navy
    C_RED = RGBColor(163, 29, 35)         # Symbiosis Accent Red
    C_BLUE_LIGHT = RGBColor(238, 242, 246)# Soft Card Background
    C_DARK = RGBColor(30, 41, 59)         # Charcoal Text
    C_MUTED = RGBColor(100, 116, 139)     # Subtitle Muted Gray
    C_WHITE = RGBColor(255, 255, 255)
    C_BORDER = RGBColor(218, 226, 237)    # Clean Border

    def add_slide_header(slide, slide_num, title_text, category="DATA STRUCTURES - LAB PROJECT"):
        # Category Tag
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.35))
        tf_c = tb_cat.text_frame
        p_c = tf_c.paragraphs[0]
        p_c.text = f"{slide_num}.  {category.upper()}"
        p_c.font.size = Pt(10)
        p_c.font.bold = True
        p_c.font.color.rgb = C_RED

        # Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.733), Inches(0.7))
        tf_t = tb_title.text_frame
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = C_NAVY

        # Accent Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(0.03))
        line.fill.solid()
        line.fill.fore_color.rgb = C_BORDER
        line.line.color.rgb = C_BORDER

        # Footer
        tb_foot = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
        tf_f = tb_foot.text_frame
        p_f = tf_f.paragraphs[0]
        p_f.text = f"CampusFind  |  Avignyah Jakku (25070721024) & Meghana (25070721002)  |  SIT Hyderabad  |  Slide {slide_num} of 6"
        p_f.font.size = Pt(9)
        p_f.font.color.rgb = C_MUTED

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)

    # University Header Tag
    tb_u = s1.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.733), Inches(0.9))
    tf_u = tb_u.text_frame
    p_u1 = tf_u.paragraphs[0]
    p_u1.text = "SYMBIOSIS INSTITUTE OF TECHNOLOGY, HYDERABAD"
    p_u1.font.size = Pt(14)
    p_u1.font.bold = True
    p_u1.font.color.rgb = C_RED

    p_u2 = tf_u.add_paragraph()
    p_u2.text = "Constituent of Symbiosis International (Deemed University), Pune  •  NAAC 'A++' Accredited"
    p_u2.font.size = Pt(10.5)
    p_u2.font.color.rgb = C_MUTED

    # Main Title Card
    c_title = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.733), Inches(2.2))
    c_title.fill.solid()
    c_title.fill.fore_color.rgb = C_NAVY
    c_title.line.color.rgb = C_NAVY
    tf_t = c_title.text_frame
    tf_t.margin_left = Inches(0.4)
    tf_t.margin_top = Inches(0.25)

    p_t1 = tf_t.paragraphs[0]
    p_t1.text = "PROJECT PROPOSAL & LAB DEFENSE"
    p_t1.font.size = Pt(11)
    p_t1.font.bold = True
    p_t1.font.color.rgb = RGBColor(254, 202, 202)

    p_t2 = tf_t.add_paragraph()
    p_t2.text = "CampusFind: Smart University Lost & Found System"
    p_t2.font.size = Pt(26)
    p_t2.font.bold = True
    p_t2.font.color.rgb = C_WHITE

    p_t3 = tf_t.add_paragraph()
    p_t3.text = "In-Memory Binary Search Tree (BST) Indexing & Multi-Factor Heuristic Matching"
    p_t3.font.size = Pt(13)
    p_t3.font.color.rgb = RGBColor(203, 213, 225)

    # Team Card
    card_team = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.2), Inches(6.0), Inches(2.4))
    card_team.fill.solid()
    card_team.fill.fore_color.rgb = C_BLUE_LIGHT
    card_team.line.color.rgb = C_BORDER
    tf_tm = card_team.text_frame
    tf_tm.margin_left = Inches(0.3)
    tf_tm.margin_top = Inches(0.2)

    p_th = tf_tm.paragraphs[0]
    p_th.text = "SUBMITTED BY:"
    p_th.font.size = Pt(11)
    p_th.font.bold = True
    p_th.font.color.rgb = C_RED

    p_tm1 = tf_tm.add_paragraph()
    p_tm1.text = "• Avignyah Jakku   |  PRN: 25070721024  (BTech CSE)"
    p_tm1.font.size = Pt(12)
    p_tm1.font.bold = True
    p_tm1.font.color.rgb = C_DARK

    p_tm2 = tf_tm.add_paragraph()
    p_tm2.text = "• Meghana                 |  PRN: 25070721002  (BTech CSE)"
    p_tm2.font.size = Pt(12)
    p_tm2.font.bold = True
    p_tm2.font.color.rgb = C_DARK

    p_tm3 = tf_tm.add_paragraph()
    p_tm3.text = "Department of Computer Science & Engineering"
    p_tm3.font.size = Pt(10.5)
    p_tm3.font.color.rgb = C_MUTED

    # Guide Card
    card_guide = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(4.2), Inches(5.433), Inches(2.4))
    card_guide.fill.solid()
    card_guide.fill.fore_color.rgb = C_BLUE_LIGHT
    card_guide.line.color.rgb = C_BORDER
    tf_g = card_guide.text_frame
    tf_g.margin_left = Inches(0.3)
    tf_g.margin_top = Inches(0.2)

    p_gh = tf_g.paragraphs[0]
    p_gh.text = "PROJECT GUIDANCE:"
    p_gh.font.size = Pt(11)
    p_gh.font.bold = True
    p_gh.font.color.rgb = C_NAVY

    p_g1 = tf_g.add_paragraph()
    p_g1.text = "Dr. Rakesh Salakapuri"
    p_g1.font.size = Pt(13)
    p_g1.font.bold = True
    p_g1.font.color.rgb = C_RED

    p_g2 = tf_g.add_paragraph()
    p_g2.text = "Professor & Project Guide, Dept. of CSE"
    p_g2.font.size = Pt(11)
    p_g2.font.color.rgb = C_DARK

    p_g3 = tf_g.add_paragraph()
    p_g3.text = "Under the Directorate of Dr. Rajinikanth Aluvalu, Director, SIT-H"
    p_g3.font.size = Pt(10)
    p_g3.font.color.rgb = C_MUTED

    # =========================================================================
    # SLIDE 2: PROBLEM STATEMENT & MOTIVATION
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_header(s2, 2, "Problem Statement & Motivation", "Context & Rationale")

    cards_s2 = [
        ("The Current Campus Challenge", C_RED, [
            "Scattered Information: Reports spread across noisy WhatsApp chats and paper registers.",
            "Vulnerability to Theft: Publicly posted item photos and details invite dishonest claims.",
            "Manual & Slow: Students and guards spend days manually searching custody desks.",
            "Zero Algorithmic Search: No indexing or automatic correlation between lost and found items."
        ]),
        ("The CampusFind Solution", C_NAVY, [
            "Centralized Digital Repository: One unified portal for all students, staff, and security desks.",
            "Anti-Fraud Challenge Verification: Secret hint questions ensure only true owners can claim items.",
            "In-Memory BST Indexing: Sub-millisecond item lookups with visual step-by-step path tracing.",
            "Automated Heuristic Matching: Correlates reports using category, keyword, location, and date."
        ])
    ]

    for idx, (title, color, points) in enumerate(cards_s2):
        x = Inches(0.8 + idx * 6.0)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(5.733), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = C_BLUE_LIGHT
        card.line.color.rgb = color
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.3)
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = color

        for pt in points:
            p_pt = tf.add_paragraph()
            p_pt.text = "\n• " + pt
            p_pt.font.size = Pt(11.5)
            p_pt.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 3: SYSTEM ARCHITECTURE & 4-STEP WORKFLOW
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_header(s3, 3, "System Architecture & End-to-End Workflow", "System Design")

    # Workflow Steps
    steps = [
        ("Step 1: Report Item", "Log lost/found item with title, campus zone, photo, and a secret ownership hint question.", C_NAVY),
        ("Step 2: BST Indexing", "Item record is saved to DB and its numerical ID is indexed into an in-memory BST in O(log n).", C_RED),
        ("Step 3: Auto-Matching", "Multi-factor algorithm evaluates similarity. Scores >= 75% trigger high-confidence matches.", C_NAVY),
        ("Step 4: Claim & Verify", "Claimant answers secret hint question; Administrator verifies and approves physical release.", C_RED)
    ]

    for idx, (s_title, s_desc, color) in enumerate(steps):
        top_y = Inches(1.7 + idx * 1.25)
        bar = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_y, Inches(11.733), Inches(1.05))
        bar.fill.solid()
        bar.fill.fore_color.rgb = C_BLUE_LIGHT
        bar.line.color.rgb = color
        bar.line.width = Pt(1.5)

        tf = bar.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.12)
        p = tf.paragraphs[0]
        p.text = s_title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color

        p_desc = tf.add_paragraph()
        p_desc.text = s_desc
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 4: CORE DSA — IN-MEMORY BINARY SEARCH TREE
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_header(s4, 4, "Applied DSA: In-Memory Binary Search Tree (BST)", "Data Structures Core")

    # Left: Theoretical concept
    c_bst_l = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8))
    c_bst_l.fill.solid()
    c_bst_l.fill.fore_color.rgb = C_BLUE_LIGHT
    c_bst_l.line.color.rgb = C_NAVY
    tf_l = c_bst_l.text_frame
    tf_l.margin_left = Inches(0.3)
    tf_l.margin_top = Inches(0.25)

    p = tf_l.paragraphs[0]
    p.text = "Why an In-Memory BST?"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    bst_bullets = [
        "Enforces BST Invariant: LeftSubtree(N) < Key(N) < RightSubtree(N) across all items.",
        "Sub-Millisecond Speed: Operates directly in RAM, avoiding disk I/O bottlenecks.",
        "Explainable Search Path: Logs exact traversal route [Root -> ... -> Target] directly to UI.",
        "In-Order Traversal: Left-Root-Right produces sorted inventory array in linear O(n) time.",
        "3-Case Deletion: Cleans resolved items (leaf, 1-child bypass, 2-children successor replacement)."
    ]
    for b in bst_bullets:
        p_b = tf_l.add_paragraph()
        p_b.text = "\n• " + b
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = C_DARK

    # Right: Complexity Table
    t_shape = s4.shapes.add_table(5, 3, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.8))
    t_bst = t_shape.table
    t_bst.columns[0].width = Inches(2.2)
    t_bst.columns[1].width = Inches(1.7)
    t_bst.columns[2].width = Inches(1.833)

    h_bst = ["BST Operation", "Average Time", "Worst Case"]
    for i, h in enumerate(h_bst):
        cell = t_bst.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_WHITE

    r_bst = [
        ["Key Search", "O(log₂ n)", "O(n)"],
        ["Key Insertion", "O(log₂ n)", "O(n)"],
        ["Three-Case Deletion", "O(log₂ n)", "O(n)"],
        ["In-Order Traversal", "O(n)", "O(n)"]
    ]
    for r_idx, r_data in enumerate(r_bst):
        for c_idx, val in enumerate(r_data):
            cell = t_bst.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_BLUE_LIGHT if r_idx % 2 == 1 else C_WHITE
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(11)
            p.font.color.rgb = C_DARK
            if c_idx > 0:
                p.font.bold = True

    # =========================================================================
    # SLIDE 5: HEURISTIC MATCHING & SECURITY VERIFICATION
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_header(s5, 5, "Multi-Factor Matching Engine & Anti-Fraud Security", "Algorithmic Logic")

    # Formula Box
    f_box = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.733), Inches(0.8))
    f_box.fill.solid()
    f_box.fill.fore_color.rgb = C_NAVY
    tf_f = f_box.text_frame
    tf_f.margin_top = Inches(0.12)
    p_f = tf_f.paragraphs[0]
    p_f.text = "MATCHING SCORE = 30% Category + 35% Keyword Overlap + 20% Location + 15% Date Plausibility"
    p_f.font.size = Pt(12)
    p_f.font.bold = True
    p_f.font.color.rgb = RGBColor(254, 202, 202)
    p_f.alignment = PP_ALIGN.CENTER

    # 4 Factor Cards
    factors = [
        ("Category Match (30%)", "30.0 pts for exact taxonomic match; 0.0 pts for mismatch (Electronics vs Bottle).", C_NAVY),
        ("Text Overlap (35%)", "Jaccard similarity on non-stopword tokens ('dell', 'laptop') boosted by title intersection.", C_RED),
        ("Campus Location (20%)", "20.0 pts for exact campus landmark match; 15.0 pts for shared campus zone.", C_NAVY),
        ("Date Plausibility (15%)", "15.0 pts if found within same day; penalizes chronologically impossible claims.", C_RED)
    ]

    for idx, (f_title, f_desc, color) in enumerate(factors):
        row = idx // 2
        col = idx % 2
        x = Inches(0.8 + col * 6.0)
        y = Inches(2.7 + row * 2.0)

        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.733), Inches(1.75))
        card.fill.solid()
        card.fill.fore_color.rgb = C_BLUE_LIGHT
        card.line.color.rgb = color
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.18)
        p = tf.paragraphs[0]
        p.text = f_title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color

        p_desc = tf.add_paragraph()
        p_desc.text = "\n" + f_desc
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 6: LIVE DEMO HIGHLIGHTS & CONCLUSION
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_header(s6, 6, "Live System Demonstration & Project Summary", "Technical Defense")

    # Demo Points Card
    c_demo = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.733), Inches(3.3))
    c_demo.fill.solid()
    c_demo.fill.fore_color.rgb = C_BLUE_LIGHT
    c_demo.line.color.rgb = C_NAVY
    c_demo.line.width = Pt(1.5)

    tf_d = c_demo.text_frame
    tf_d.margin_left = Inches(0.4)
    tf_d.margin_top = Inches(0.25)

    p_dh = tf_d.paragraphs[0]
    p_dh.text = "CORE FEATURES TO OBSERVE IN THE LIVE DEMO:"
    p_dh.font.size = Pt(13)
    p_dh.font.bold = True
    p_dh.font.color.rgb = C_NAVY

    demos = [
        "1. Dynamic BST Visualizer (/dsa/visualizer): Real-time SVG tree rendering with green laser search path animation.",
        "2. Heuristic Matching Hub (/matches): Automated item pairing with percentage scores and transparent factor breakdowns.",
        "3. Anti-Fraud Reporting (/report): Confidential owner challenge questions that prevent false claims.",
        "4. Administrative Portal (/admin): Approving authentic claims, rejecting fakes, and viewing immutable audit trails.",
        "5. Engineering Quality: Dual MySQL/SQLite persistence verified across 51/51 automated unit and integration tests."
    ]
    for d in demos:
        p_pt = tf_d.add_paragraph()
        p_pt.text = "• " + d
        p_pt.font.size = Pt(11)
        p_pt.font.color.rgb = C_DARK

    # Live Demo & Q&A Banner
    b_qa = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.2), Inches(11.733), Inches(1.5))
    b_qa.fill.solid()
    b_qa.fill.fore_color.rgb = C_NAVY
    tf_qa = b_qa.text_frame
    tf_qa.margin_top = Inches(0.2)

    p_q1 = tf_qa.paragraphs[0]
    p_q1.text = "READY FOR LIVE SYSTEM DEMONSTRATION & Q&A"
    p_q1.font.size = Pt(16)
    p_q1.font.bold = True
    p_q1.font.color.rgb = C_WHITE
    p_q1.alignment = PP_ALIGN.CENTER

    p_q2 = tf_qa.add_paragraph()
    p_q2.text = "Avignyah Jakku (25070721024)   &   Meghana (25070721002)"
    p_q2.font.size = Pt(12)
    p_q2.font.bold = True
    p_q2.font.color.rgb = RGBColor(254, 202, 202)
    p_q2.alignment = PP_ALIGN.CENTER

    p_q3 = tf_qa.add_paragraph()
    p_q3.text = "Department of Computer Science & Engineering  |  SIT Hyderabad"
    p_q3.font.size = Pt(10)
    p_q3.font.color.rgb = RGBColor(203, 213, 225)
    p_q3.alignment = PP_ALIGN.CENTER

    prs.save(output_filename)
    print(f"Simple 6-slide presentation saved successfully to: {output_filename}")

if __name__ == '__main__':
    create_simple_deck()
